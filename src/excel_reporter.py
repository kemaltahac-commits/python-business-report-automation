from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from analyzer import AnalysisResult


def _format_sheet(writer: pd.ExcelWriter, sheet_name: str, currency_columns: tuple[str, ...] = ()) -> None:
    worksheet = writer.sheets[sheet_name]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        worksheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 28)
    headers = [cell.value for cell in worksheet[1]]
    for name in currency_columns:
        if name in headers:
            index = headers.index(name) + 1
            for cell in list(worksheet.columns)[index - 1][1:]:
                cell.number_format = '$#,##0.00'


def generate_excel_report(result: AnalysisResult, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    summary = pd.DataFrame({
        "KPI": ["Total Sales", "Total Orders", "Total Quantity", "Average Order Value", "Best-Performing Region", "Reporting Period"],
        "Value": [result.total_sales, result.total_orders, result.total_quantity, result.average_order_value,
                  result.best_region, f"{result.reporting_start:%Y-%m-%d} to {result.reporting_end:%Y-%m-%d}"],
    })
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        result.sales_by_product.to_excel(writer, sheet_name="Sales by Product", index=False)
        result.sales_by_region.to_excel(writer, sheet_name="Sales by Region", index=False)
        result.sales_by_customer.to_excel(writer, sheet_name="Sales by Customer", index=False)
        result.quantity_by_product.to_excel(writer, sheet_name="Product Quantity", index=False)
        result.monthly_sales.to_excel(writer, sheet_name="Monthly Sales", index=False)
        _format_sheet(writer, "Summary")
        for name in ("Sales by Product", "Sales by Region", "Sales by Customer", "Monthly Sales"):
            _format_sheet(writer, name, ("Sales",))
        _format_sheet(writer, "Product Quantity")
        summary_ws = writer.sheets["Summary"]
        for row in (2, 5):
            summary_ws.cell(row, 2).number_format = '$#,##0.00'
